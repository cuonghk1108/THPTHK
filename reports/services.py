
def get_natural_class_sort_key(score_item):
    c = score_item.school_class
    order = getattr(c, 'class_order', None)
    if order is not None:
        return (c.grade or 0, order, c.name)
    m = re.search(r'\d+', c.name.replace(str(c.grade or ''), '', 1))
    num = int(m.group(0)) if m else 999
    return (c.grade or 0, num, c.name)

# -*- coding: utf-8 -*-
"""
Dịch vụ xuất báo cáo định dạng Excel (.xlsx) và PDF (.pdf) chuẩn hóa sử dụng Font Times New Roman
dành cho hệ thống quản lý học sinh HK Student Hub – Trường THPT Hương Khê, Hà Tĩnh.
"""
import io
from datetime import datetime
from decimal import Decimal
from django.core.exceptions import PermissionDenied
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable

from competition.models import CompetitionPeriod, ClassCompetitionScore
from competition.services import get_competition_class_scope, get_period_status, get_valid_aggregate_class_ids
from conduct.models import ViolationRecord
from .models import ReportExportHistory
from .pdf_utils import (
    setup_times_new_roman_fonts,
    NumberedCanvas,
    get_standard_pdf_styles,
    safe_filename,
)


def export_competition_excel(period_id, exported_by=None):
    """Xuất bảng tổng hợp thi đua Chi đoàn ra tệp Excel (.xlsx) tiêu chuẩn"""
    period = period_id if isinstance(period_id, CompetitionPeriod) else CompetitionPeriod.objects.get(id=period_id)
    if exported_by and not period.is_published and not (exported_by.is_admin or exported_by.is_bgh or exported_by.is_doan_truong):
        raise PermissionDenied('Kỳ thi đua chưa được công bố.')
    # Đảm bảo tính lại điểm thi đua thực tế mới nhất trước khi xuất file
    from competition.services import recalculate_period_scores
    recalculate_period_scores(period, force=True)

    scores_qs = ClassCompetitionScore.objects.filter(competition_period=period).select_related('school_class')
    if exported_by and not (exported_by.is_admin or exported_by.is_bgh or exported_by.is_doan_truong):
        scores_qs = scores_qs.filter(school_class__in=get_competition_class_scope(exported_by, period.academic_year))
    valid_ids = get_valid_aggregate_class_ids(period)
    if valid_ids is not None:
        scores_qs = scores_qs.filter(school_class_id__in=valid_ids)
    
    # Sắp xếp các lớp từ 10A1-10A22, 11A1-11A21, 12A1-12A22 theo đúng yêu cầu
    scores = sorted(list(scores_qs), key=get_natural_class_sort_key)
    period_state = get_period_status(period)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"ThiDua_Tuan_{period.week_number or 'Ky'}"[:30]

    # Styles
    title_font = Font(name='Times New Roman', size=14, bold=True, color='1E3A8A')
    sub_font = Font(name='Times New Roman', size=11, bold=True, color='374151')
    header_font = Font(name='Times New Roman', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # Headers
    ws.merge_cells('A1:I1')
    ws['A1'] = "SỞ GIÁO DỤC VÀ ĐÀO TẠO HÀ TĨNH - TRƯỜNG THPT HƯƠNG KHÊ"
    ws['A1'].font = sub_font
    ws['A1'].alignment = Alignment(horizontal='left')

    ws.merge_cells('A2:I2')
    ws['A2'] = f"BẢNG TỔNG HỢP XẾP HẠNG THI ĐUA CHI ĐOÀN – {period.name.upper()}"
    ws['A2'].font = title_font
    ws['A2'].alignment = Alignment(horizontal='center')

    ws.merge_cells('A3:I3')
    now_str = timezone.now().strftime('%d/%m/%Y %H:%M')
    creator_str = exported_by.full_name or exported_by.username if exported_by else 'Ban Thi đua Đoàn trường'
    ws['A3'] = f"Trạng thái: {period_state['display']} | Thời gian áp dụng: {period.start_date.strftime('%d/%m/%Y')} đến {period.end_date.strftime('%d/%m/%Y')} | Xuất lúc: {now_str} | Người lập: {creator_str}"
    ws['A3'].font = Font(name='Times New Roman', size=9, italic=True, color='6B7280')
    ws['A3'].alignment = Alignment(horizontal='center')

    headers = [
        "Thứ hạng", "Lớp / Chi đoàn", "Khối", "Điểm gốc", "Điểm cộng", 
        "Điểm trừ vi phạm", "Tổng điểm thi đua", "Số lỗi vi phạm", "Biến động hạng"
    ]
    
    ws.append([])
    ws.append(headers)
    header_row = 5
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    row_idx = 6
    for item in scores:
        trend = "0"
        if item.rank_change is not None:
            if item.rank_change > 0:
                trend = f"+{item.rank_change} ▲"
            elif item.rank_change < 0:
                trend = f"{item.rank_change} ▼"
        else:
            trend = "-- (Chưa có dữ liệu so sánh)"

        row_data = [
            item.rank,
            item.school_class.name,
            f"Khối {item.school_class.grade}",
            float(item.base_points),
            float(item.bonus_points),
            float(item.penalty_points),
            float(item.total_score),
            item.verified_violations_count,
            trend
        ]
        ws.append(row_data)

        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=c_idx)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center' if c_idx in [1, 3, 8, 9] else ('right' if c_idx in [4, 5, 6, 7] else 'left'))
            if c_idx == 7:
                cell.font = Font(name='Times New Roman', size=10, bold=True, color='1E3A8A')
        row_idx += 1

    ws.append([])
    footer_row = ws.max_row
    ws.append(['HK Student Hub – Trường THPT Hương Khê • code by cuongdev1108'])
    ws.cell(row=ws.max_row, column=1).font = Font(name='Times New Roman', size=9, italic=True, color='64748B')

    column_widths = [12, 18, 12, 14, 14, 18, 18, 16, 16]
    for i, width in enumerate(column_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    ReportExportHistory.objects.create(
        report_type=ReportExportHistory.REPORT_COMPETITION,
        file_format=ReportExportHistory.FORMAT_EXCEL,
        title=f"Báo cáo thi đua {period.name}",
        parameters=f'{{"period_id": {period.id}}}',
        exported_by=exported_by
    )

    return output.getvalue()


def export_competition_pdf(period_id, exported_by=None):
    """
    Xuất báo cáo thi đua Chi đoàn dạng PDF chuẩn hóa sử dụng Font Times New Roman 100%.
    Hỗ trợ đầy đủ tiếng Việt có dấu, header/footer đánh số trang "Trang X / Y".
    """
    period = period_id if isinstance(period_id, CompetitionPeriod) else CompetitionPeriod.objects.get(id=period_id)
    if exported_by and not period.is_published and not (exported_by.is_admin or exported_by.is_bgh or exported_by.is_doan_truong):
        raise PermissionDenied('Kỳ thi đua chưa được công bố.')
    # Đảm bảo tính lại điểm thi đua thực tế mới nhất trước khi xuất file
    from competition.services import recalculate_period_scores
    recalculate_period_scores(period, force=True)

    scores_qs = ClassCompetitionScore.objects.filter(competition_period=period).select_related('school_class')
    if exported_by and not (exported_by.is_admin or exported_by.is_bgh or exported_by.is_doan_truong):
        scores_qs = scores_qs.filter(school_class__in=get_competition_class_scope(exported_by, period.academic_year))
    valid_ids = get_valid_aggregate_class_ids(period)
    if valid_ids is not None:
        scores_qs = scores_qs.filter(school_class_id__in=valid_ids)
    
    # Sắp xếp các lớp từ 10A1-10A22, 11A1-11A21, 12A1-12A22 theo đúng yêu cầu
    scores = sorted(list(scores_qs), key=get_natural_class_sort_key)
    period_state = get_period_status(period)

    setup_times_new_roman_fonts()
    styles = get_standard_pdf_styles()

    buffer = io.BytesIO()
    # A4 Portrait: 210 x 297 mm (~595 x 842 pt), lề 15mm (~42.5pt)
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=42,
        rightMargin=42,
        topMargin=45,
        bottomMargin=45
    )

    elements = []

    # 1. Tiêu đề cơ quan ban hành (Chuẩn hành chính văn bản giáo dục)
    header_data = [
        [
            Paragraph("SỞ GD&ĐT HÀ TĨNH<br/><b>TRƯỜNG THPT HƯƠNG KHÊ</b><br/><b>BAN THI ĐUA – ĐOÀN TRƯỜNG</b>", styles['OrgHeader']),
            Paragraph("<b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br/><b><u>Độc lập – Tự do – Hạnh phúc</u></b>", styles['OrgHeaderRight'])
        ]
    ]
    header_table = Table(header_data, colWidths=[260, 250])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 14))

    # 2. Tiêu đề báo cáo
    elements.append(Paragraph(f"BẢNG TỔNG HỢP XẾP HẠNG THI ĐUA CHI ĐOÀN", styles['DocTitle']))
    elements.append(Paragraph(f"<b>{period.name.upper()}</b>", styles['DocTitle']))
    
    now_str = timezone.now().strftime('%d/%m/%Y %H:%M')
    creator_name = exported_by.full_name or exported_by.username if exported_by else 'Ban Thi đua Đoàn trường'
    period_str = f"Trạng thái: {period_state['display']}<br/>Thời gian áp dụng: Từ {period.start_date.strftime('%d/%m/%Y')} đến {period.end_date.strftime('%d/%m/%Y')}"
    elements.append(Paragraph(f"{period_str}<br/><i>(Xuất lúc: {now_str} – Người lập: {creator_name})</i>", styles['DocSubtitle']))
    elements.append(Spacer(1, 6))

    # 3. Bảng dữ liệu xếp hạng thi đua
    col_widths = [35, 80, 50, 60, 60, 60, 75, 45, 45]
    table_headers = [
        Paragraph("<b>Hạng</b>", styles['TableHead']),
        Paragraph("<b>Chi đoàn</b>", styles['TableHead']),
        Paragraph("<b>Khối</b>", styles['TableHead']),
        Paragraph("<b>Điểm gốc</b>", styles['TableHead']),
        Paragraph("<b>Điểm cộng</b>", styles['TableHead']),
        Paragraph("<b>Điểm trừ</b>", styles['TableHead']),
        Paragraph("<b>Tổng điểm</b>", styles['TableHead']),
        Paragraph("<b>Số lỗi</b>", styles['TableHead']),
        Paragraph("<b>Biến động</b>", styles['TableHead']),
    ]
    data = [table_headers]

    for sc in scores:
        trend_str = "0"
        if sc.rank_change is not None:
            if sc.rank_change > 0:
                trend_str = f"+{sc.rank_change}"
            elif sc.rank_change < 0:
                trend_str = f"{sc.rank_change}"
        else:
            trend_str = "--"

        row = [
            Paragraph(str(sc.rank), styles['TableCellCenter']),
            Paragraph(f"<b>{sc.school_class.name}</b>", styles['TableCellBold']),
            Paragraph(f"Khối {sc.school_class.grade}", styles['TableCellCenter']),
            Paragraph(f"{sc.base_points:.2f}đ", styles['TableCellCenter']),
            Paragraph(f"+{sc.bonus_points:.2f}đ" if sc.bonus_points > 0 else "0.00đ", styles['TableCellCenter']),
            Paragraph(f"-{sc.penalty_points:.2f}đ" if sc.penalty_points > 0 else "0.00đ", styles['TableCellCenter']),
            Paragraph(f"<b>{sc.total_score:.2f}đ</b>", styles['TableCellBold']),
            Paragraph(str(sc.verified_violations_count), styles['TableCellCenter']),
            Paragraph(trend_str, styles['TableCellCenter']),
        ]
        data.append(row)

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#163A70')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F8FAFC'), colors.white]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 20))

    # 4. Khu vực chữ ký xác nhận
    sig_data = [
        [
            "",
            Paragraph("<i>Hương Khê, ngày ..... tháng ..... năm 2026</i>", styles['SignatureDate'])
        ],
        [
            Paragraph("<b>TM. ĐOÀN TRƯỜNG</b><br/><b>BÍ THƯ</b>", styles['SignatureTitle']),
            Paragraph("<b>HIỆU TRƯỞNG</b><br/><i>(Phê duyệt bảng thi đua)</i>", styles['SignatureTitle'])
        ],
        [
            Paragraph("<br/><br/><br/><i>(Ký, ghi rõ họ tên)</i>", styles['SignatureDate']),
            Paragraph("<br/><br/><br/><i>(Ký tên và đóng dấu)</i>", styles['SignatureDate'])
        ]
    ]
    sig_table = Table(sig_data, colWidths=[255, 255])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    elements.append(KeepTogether(sig_table))

    doc.build(elements, canvasmaker=NumberedCanvas)
    buffer.seek(0)

    ReportExportHistory.objects.create(
        report_type=ReportExportHistory.REPORT_COMPETITION,
        file_format=ReportExportHistory.FORMAT_PDF,
        title=f"Báo cáo PDF thi đua {period.name}",
        parameters=f'{{"period_id": {period.id}}}',
        exported_by=exported_by
    )

    return buffer.getvalue()


def export_violations_excel(filters=None, exported_by=None):
    """Xuất danh sách vi phạm nền nếp ra Excel với bộ lọc"""
    queryset = ViolationRecord.objects.select_related('school_class', 'student', 'rule', 'reporter').all().order_by('-occurred_at')
    queryset = _scope_violation_queryset(queryset, exported_by)
    if filters:
        if filters.get('class_id'):
            queryset = queryset.filter(school_class_id=filters['class_id'])
        if filters.get('status'):
            queryset = queryset.filter(status=filters['status'])
        if filters.get('start_date'):
            queryset = queryset.filter(occurred_at__date__gte=filters['start_date'])
        if filters.get('end_date'):
            queryset = queryset.filter(occurred_at__date__lte=filters['end_date'])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DS Vi pham nen nep"

    headers = ["Mã bản ghi", "Thời gian", "Lớp", "Học sinh", "Lỗi vi phạm", "Vị trí", "Điểm trừ", "Trạng thái", "Người ghi nhận"]
    ws.append(["TRƯỜNG THPT HƯƠNG KHÊ - BÁO CÁO VI PHẠM NỀN NẾP HỌC SINH"])
    ws.append([f"Xuất ngày: {timezone.now().strftime('%d/%m/%Y %H:%M')}"])
    ws.append([])
    ws.append(headers)

    header_font = Font(name='Times New Roman', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for rec in queryset:
        ws.append([
            rec.record_code,
            rec.occurred_at.strftime('%d/%m/%Y %H:%M') if rec.occurred_at else '',
            rec.school_class.name if rec.school_class else '',
            rec.student_name_snapshot or (rec.student.full_name if rec.student else "Tập thể lớp"),
            rec.rule.name if rec.rule else '',
            rec.location or '',
            float(rec.points_deducted),
            rec.get_status_display(),
            rec.reporter.full_name or rec.reporter.username if rec.reporter else 'Hệ thống'
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    ReportExportHistory.objects.create(
        report_type=ReportExportHistory.REPORT_VIOLATIONS,
        file_format=ReportExportHistory.FORMAT_EXCEL,
        title="Danh sách vi phạm nền nếp",
        parameters=str(filters or {}),
        exported_by=exported_by
    )

    return output.getvalue()


def export_violations_pdf(filters=None, exported_by=None):
    """
    Xuất danh sách vi phạm nền nếp ra PDF khổ ngang A4 chuẩn hóa sử dụng Font Times New Roman.
    Bảo đảm hiển thị đầy đủ tiếng Việt, lặp lại tiêu đề bảng khi sang trang và hiển thị số trang.
    """
    queryset = ViolationRecord.objects.select_related('school_class', 'student', 'rule', 'reporter').all().order_by('-occurred_at')
    queryset = _scope_violation_queryset(queryset, exported_by)
    filter_desc = []
    if filters:
        if filters.get('class_id'):
            queryset = queryset.filter(school_class_id=filters['class_id'])
            from conduct.models import SchoolClass
            cls_obj = SchoolClass.objects.filter(id=filters['class_id']).first()
            if cls_obj:
                filter_desc.append(f"Lớp: {cls_obj.name}")
        if filters.get('status'):
            queryset = queryset.filter(status=filters['status'])
            filter_desc.append(f"Trạng thái: {filters['status']}")
        if filters.get('start_date'):
            queryset = queryset.filter(occurred_at__date__gte=filters['start_date'])
            filter_desc.append(f"Từ ngày: {filters['start_date']}")
        if filters.get('end_date'):
            queryset = queryset.filter(occurred_at__date__lte=filters['end_date'])
            filter_desc.append(f"Đến ngày: {filters['end_date']}")

    setup_times_new_roman_fonts()
    styles = get_standard_pdf_styles()

    buffer = io.BytesIO()
    # A4 Landscape: 297 x 210 mm (~842 x 595 pt), lề 15mm (~42.5pt)
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=40,
        rightMargin=40,
        topMargin=42,
        bottomMargin=42
    )

    elements = []

    # 1. Cơ quan ban hành
    header_data = [
        [
            Paragraph("SỞ GD&ĐT HÀ TĨNH<br/><b>TRƯỜNG THPT HƯƠNG KHÊ</b><br/><b>BAN THI ĐUA – ĐOÀN TRƯỜNG</b>", styles['OrgHeader']),
            Paragraph("<b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br/><b><u>Độc lập – Tự do – Hạnh phúc</u></b>", styles['OrgHeaderRight'])
        ]
    ]
    header_table = Table(header_data, colWidths=[380, 380])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 12))

    # 2. Tiêu đề
    elements.append(Paragraph("BÁO CÁO TỔNG HỢP DANH SÁCH VI PHẠM NỀN NẾP", styles['DocTitle']))
    now_str = timezone.now().strftime('%d/%m/%Y %H:%M')
    creator_name = exported_by.full_name or exported_by.username if exported_by else 'Ban Quản lý nền nếp'
    subtitle_parts = [f"Xuất lúc: {now_str}", f"Người lập: {creator_name}"]
    if filter_desc:
        subtitle_parts.append(f"Phạm vi: {', '.join(filter_desc)}")
    elements.append(Paragraph(" | ".join(subtitle_parts), styles['DocSubtitle']))
    elements.append(Spacer(1, 8))

    # 3. Thống kê tóm tắt
    total_violations = queryset.count()
    total_deducted = sum([rec.points_deducted for rec in queryset])
    summary_text = (
        f"<b>Tổng số bản ghi vi phạm:</b> {total_violations} trường hợp   |   "
        f"<b>Tổng điểm trừ thi đua:</b> -{total_deducted:.2f} điểm"
    )
    elements.append(Paragraph(summary_text, styles['Body']))
    elements.append(Spacer(1, 8))

    # 4. Bảng dữ liệu vi phạm (Khổ ngang 760pt rộng rãi)
    col_widths = [30, 85, 75, 45, 115, 200, 50, 75, 85]
    table_headers = [
        Paragraph("<b>STT</b>", styles['TableHead']),
        Paragraph("<b>Mã biên bản</b>", styles['TableHead']),
        Paragraph("<b>Thời gian</b>", styles['TableHead']),
        Paragraph("<b>Lớp</b>", styles['TableHead']),
        Paragraph("<b>Học sinh vi phạm</b>", styles['TableHead']),
        Paragraph("<b>Hành vi vi phạm</b>", styles['TableHead']),
        Paragraph("<b>Điểm trừ</b>", styles['TableHead']),
        Paragraph("<b>Trạng thái</b>", styles['TableHead']),
        Paragraph("<b>Người ghi nhận</b>", styles['TableHead']),
    ]
    data = [table_headers]

    for idx, rec in enumerate(queryset, start=1):
        time_str = timezone.localtime(rec.occurred_at).strftime('%d/%m/%Y %H:%M') if rec.occurred_at else ''
        student_name = rec.student_name_snapshot or (rec.student.full_name if rec.student else "Tập thể lớp")
        status_color = '#059669' if rec.status == ViolationRecord.STATUS_VERIFIED else '#D97706'
        status_html = f"<font color='{status_color}'><b>{rec.get_status_display()}</b></font>"

        row = [
            Paragraph(str(idx), styles['TableCellCenter']),
            Paragraph(f"<b>{rec.record_code}</b>", styles['TableCell']),
            Paragraph(time_str, styles['TableCellCenter']),
            Paragraph(rec.school_class.name if rec.school_class else '', styles['TableCellCenter']),
            Paragraph(f"<b>{student_name}</b>", styles['TableCell']),
            Paragraph(rec.rule.name if rec.rule else (rec.description or ''), styles['TableCell']),
            Paragraph(f"<b>-{rec.points_deducted:.2f}đ</b>", styles['TableCellCenter']),
            Paragraph(status_html, styles['TableCellCenter']),
            Paragraph(rec.reporter.full_name or rec.reporter.username if rec.reporter else 'Hệ thống', styles['TableCell']),
        ]
        data.append(row)

    if len(data) == 1:
        data.append([
            Paragraph("Không có bản ghi vi phạm nào phù hợp với bộ lọc.", styles['TableCellCenter']),
            "", "", "", "", "", "", "", ""
        ])

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#163A70')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F8FAFC'), colors.white]),
    ]
    if len(data) == 2 and len(data[1]) > 1 and data[1][1] == "":
        table_style.append(('SPAN', (0, 1), (-1, 1)))
    table.setStyle(TableStyle(table_style))
    elements.append(table)
    elements.append(Spacer(1, 20))

    # 5. Chữ ký
    sig_data = [
        [
            "",
            Paragraph("<i>Hương Khê, ngày ..... tháng ..... năm 2026</i>", styles['SignatureDate'])
        ],
        [
            Paragraph("<b>NGƯỜI LẬP BÁO CÁO</b>", styles['SignatureTitle']),
            Paragraph("<b>TM. BAN CHẤP HÀNH ĐOÀN TRƯỜNG</b>", styles['SignatureTitle'])
        ],
        [
            Paragraph("<br/><br/><br/><i>(Ký, ghi rõ họ tên)</i>", styles['SignatureDate']),
            Paragraph("<br/><br/><br/><i>(Ký tên và đóng dấu)</i>", styles['SignatureDate'])
        ]
    ]
    sig_table = Table(sig_data, colWidths=[380, 380])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    elements.append(KeepTogether(sig_table))

    doc.build(elements, canvasmaker=NumberedCanvas)
    buffer.seek(0)

    ReportExportHistory.objects.create(
        report_type=ReportExportHistory.REPORT_VIOLATIONS,
        file_format=ReportExportHistory.FORMAT_PDF,
        title="Báo cáo PDF danh sách vi phạm nền nếp",
        parameters=str(filters or {}),
        exported_by=exported_by
    )

    return buffer.getvalue()


def export_violation_detail_pdf(violation_id, exported_by=None):
    """
    Xuất Biên bản ghi nhận vi phạm nền nếp & Minh chứng trừ điểm chính thức dạng PDF.
    Sử dụng 100% Font Times New Roman, đầy đủ chữ ký các bên, bảo mật minh chứng.
    """
    violation = ViolationRecord.objects.select_related(
        'school_class', 'student', 'rule', 'rule__category', 'reporter', 'verified_by'
    ).get(id=violation_id)
    if exported_by:
        scoped = _scope_violation_queryset(ViolationRecord.objects.filter(pk=violation.pk), exported_by)
        if not scoped.exists():
            raise PermissionDenied('Bạn không có quyền xem bản ghi này.')

    setup_times_new_roman_fonts()
    styles = get_standard_pdf_styles()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=45,
        rightMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    elements = []

    # 1. Quốc hiệu / Tiêu ngữ và Tên trường
    header_data = [
        [
            Paragraph("SỞ GD&ĐT HÀ TĨNH<br/><b>TRƯỜNG THPT HƯƠNG KHÊ</b><br/><b>ĐOÀN TRƯỜNG – TỔ GIÁM THỊ</b>", styles['OrgHeader']),
            Paragraph("<b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br/><b><u>Độc lập – Tự do – Hạnh phúc</u></b>", styles['OrgHeaderRight'])
        ]
    ]
    header_table = Table(header_data, colWidths=[250, 255])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 15))

    # 2. Tiêu đề Biên bản
    elements.append(Paragraph("BIÊN BẢN GHI NHẬN VI PHẠM NỀN NẾP", styles['DocTitle']))
    elements.append(Paragraph("<b>VÀ XÁC NHẬN MINH CHỨNG TRỪ ĐIỂM THI ĐUA</b>", styles['DocTitle']))
    elements.append(Paragraph(f"Mã biên bản: <b>{violation.record_code}</b>", styles['DocSubtitle']))
    elements.append(Spacer(1, 10))

    # 3. Thông tin thời gian và địa điểm
    occurred_str = timezone.localtime(violation.occurred_at).strftime('%H giờ %M phút, ngày %d tháng %m năm %Y') if violation.occurred_at else 'N/A'
    created_str = timezone.localtime(violation.created_at).strftime('%d/%m/%Y %H:%M') if violation.created_at else 'N/A'
    location_str = violation.location or 'Khuôn viên Trường THPT Hương Khê'

    student_name = violation.student.full_name if violation.student else (violation.student_name_snapshot or f"Tập thể lớp {violation.school_class.name if violation.school_class else ''}")
    class_name = violation.school_class.name if violation.school_class else 'N/A'
    grade_str = f"Khối {violation.school_class.grade}" if violation.school_class else 'N/A'

    reporter_name = violation.reporter.full_name or violation.reporter.username if violation.reporter else 'Hệ thống'
    reporter_role = violation.reporter.get_role_display() if violation.reporter else 'Cán bộ giám thị'

    intro_p = (
        f"Hôm nay, vào lúc {occurred_str}, tại: {location_str}, "
        f"chúng tôi tiến hành lập biên bản ghi nhận hành vi vi phạm nền nếp học đường với các nội dung chi tiết dưới đây:"
    )
    elements.append(Paragraph(intro_p, styles['Body']))
    elements.append(Spacer(1, 10))

    # 4. Bảng thông tin chi tiết biên bản
    info_rows = [
        [
            Paragraph("<b>I. ĐỐI TƯỢNG VI PHẠM</b>", styles['TableCellBold']),
            ""
        ],
        [
            Paragraph("Học sinh / Tập thể:", styles['TableCell']),
            Paragraph(f"<b>{student_name}</b>", styles['TableCellBold'])
        ],
        [
            Paragraph("Chi đoàn / Lớp học:", styles['TableCell']),
            Paragraph(f"Lớp <b>{class_name}</b> ({grade_str})", styles['TableCell'])
        ],
        [
            Paragraph("Người lập biên bản:", styles['TableCell']),
            Paragraph(f"<b>{reporter_name}</b> (Vai trò: {reporter_role})", styles['TableCell'])
        ],
        [
            Paragraph("<b>II. HÀNH VI VÀ QUY ĐỊNH VI PHẠM</b>", styles['TableCellBold']),
            ""
        ],
        [
            Paragraph("Quy chế vi phạm:", styles['TableCell']),
            Paragraph(f"<b>{violation.rule.name if violation.rule else 'N/A'}</b> (Mã: {violation.rule.code if violation.rule else ''})", styles['TableCell'])
        ],
        [
            Paragraph("Nhóm hành vi:", styles['TableCell']),
            Paragraph(violation.rule.category.name if violation.rule and violation.rule.category else "Nề nếp kỷ cương chung", styles['TableCell'])
        ],
        [
            Paragraph("Mô tả chi tiết vi phạm:", styles['TableCell']),
            Paragraph(violation.description or "Không có mô tả bổ sung.", styles['TableCell'])
        ],
        [
            Paragraph("<b>III. HÌNH THỨC XỬ LÝ VÀ ĐIỂM THI ĐUA</b>", styles['TableCellBold']),
            ""
        ],
        [
            Paragraph("Điểm thi đua bị trừ:", styles['TableCell']),
            Paragraph(f"<b><font color='#DC2626'>-{violation.points_deducted:.2f} điểm</font></b>", styles['TableCellBold'])
        ],
        [
            Paragraph("Điểm thi đua lớp sau khi trừ:", styles['TableCell']),
            Paragraph(
                f"<b>{violation.score_after:.2f} điểm</b> (trước khi trừ: {violation.score_before:.2f} điểm)"
                if violation.score_after is not None and violation.score_before is not None
                else (f"<b>{violation.score_after:.2f} điểm</b>" if violation.score_after is not None else "Đang cập nhật theo chu kỳ")
                , styles['TableCell']
            )
        ],
        [
            Paragraph("Trạng thái phê duyệt:", styles['TableCell']),
            Paragraph(f"<b>{violation.get_status_display()}</b>", styles['TableCell'])
        ],
        [
            Paragraph("<b>IV. MINH CHỨNG VÀ BẢO MẬT DỮ LIỆU</b>", styles['TableCellBold']),
            ""
        ],
        [
            Paragraph("Trạng thái minh chứng:", styles['TableCell']),
            Paragraph(
                "✓ Đã đính kèm minh chứng hình ảnh (Lưu trữ an toàn trên máy chủ HK Student Hub)"
                if violation.evidence_image
                else "Không có tệp ảnh đính kèm",
                styles['TableCell']
            )
        ],
        [
            Paragraph("Ghi chú minh chứng:", styles['TableCell']),
            Paragraph(violation.evidence_notes or "Không có ghi chú thêm.", styles['TableCell'])
        ],
    ]

    detail_table = Table(info_rows, colWidths=[150, 355])
    detail_table.setStyle(TableStyle([
        ('SPAN', (0, 0), (1, 0)),
        ('SPAN', (0, 4), (1, 4)),
        ('SPAN', (0, 8), (1, 8)),
        ('SPAN', (0, 12), (1, 12)),
        ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#F1F5F9')),
        ('BACKGROUND', (0, 4), (1, 4), colors.HexColor('#F1F5F9')),
        ('BACKGROUND', (0, 8), (1, 8), colors.HexColor('#F1F5F9')),
        ('BACKGROUND', (0, 12), (1, 12), colors.HexColor('#F1F5F9')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(detail_table)
    elements.append(Spacer(1, 12))

    conclude_p = (
        "Biên bản được lập thành văn bản điện tử và lưu trữ bảo mật trên hệ thống HK Student Hub. "
        "Dữ liệu được dùng làm căn cứ chính thức để tính điểm thi đua tuần/kỳ và thông báo đến Đoàn trường, Ban Giám hiệu và Giáo viên chủ nhiệm."
    )
    elements.append(Paragraph(conclude_p, styles['BodyItalic']))
    elements.append(Spacer(1, 20))

    # 5. Khu vực chữ ký xác nhận 3 bên
    sig_data = [
        [
            Paragraph("<b>HỌC SINH / ĐẠI DIỆN LỚP</b><br/><i>(Ký và ghi rõ họ tên)</i>", styles['SignatureTitle']),
            Paragraph("<b>NGƯỜI LẬP BIÊN BẢN</b><br/><i>(Ký và ghi rõ họ tên)</i>", styles['SignatureTitle']),
            Paragraph("<b>TM. ĐOÀN TRƯỜNG / BGH</b><br/><i>(Xác nhận & lưu trữ)</i>", styles['SignatureTitle']),
        ],
        [
            Paragraph("<br/><br/><br/><br/>", styles['TableCellCenter']),
            Paragraph("<br/><br/><br/><br/>", styles['TableCellCenter']),
            Paragraph("<br/><br/><br/><br/>", styles['TableCellCenter']),
        ]
    ]
    sig_table = Table(sig_data, colWidths=[165, 175, 165])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    elements.append(KeepTogether(sig_table))

    doc.build(elements, canvasmaker=NumberedCanvas)
    buffer.seek(0)

    ReportExportHistory.objects.create(
        report_type=ReportExportHistory.REPORT_VIOLATIONS,
        file_format=ReportExportHistory.FORMAT_PDF,
        title=f"Biên bản vi phạm {violation.record_code}",
        parameters=f'{{"violation_id": {violation.id}}}',
        exported_by=exported_by
    )

    return buffer.getvalue()


def _scope_violation_queryset(queryset, exported_by):
    if not exported_by or exported_by.is_admin or exported_by.is_bgh or exported_by.is_doan_truong:
        return queryset
    allowed_classes = get_competition_class_scope(exported_by)
    queryset = queryset.filter(school_class__in=allowed_classes)
    if exported_by.is_student:
        student = getattr(exported_by, 'student_profile', None)
        return queryset.filter(student=student) if student else queryset.none()
    return queryset
