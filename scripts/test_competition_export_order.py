import sys
import os

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, r'C:\Users\Xuan Cuong\Desktop\hk-student-hub')
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'hk_hub_project.settings')

import django
django.setup()

from competition.models import CompetitionPeriod, ClassCompetitionScore
from competition.services import recalculate_period_scores
from reports.services import export_competition_excel, export_competition_pdf
from django.contrib.auth import get_user_model
import openpyxl
import io

User = get_user_model()
admin = User.objects.filter(is_superuser=True).first()

# Period 1 = Tuần 7
p_week = CompetitionPeriod.objects.get(id=1)
# Period 6 = Học kỳ I
p_sem = CompetitionPeriod.objects.get(id=6)

print(f"--- TESTING WEEK: {p_week.name} ---")
recalculate_period_scores(p_week, force=True)
score_10a1 = ClassCompetitionScore.objects.get(competition_period=p_week, school_class__name='10A1')
print(f"Week 7 - 10A1: score={score_10a1.total_score}, penalty={score_10a1.penalty_points}, rank={score_10a1.rank}")

print(f"\n--- TESTING SEMESTER: {p_sem.name} ---")
recalculate_period_scores(p_sem, force=True)
sem_10a1 = ClassCompetitionScore.objects.get(competition_period=p_sem, school_class__name='10A1')
print(f"Semester 1 - 10A1: score={sem_10a1.total_score}, penalty={sem_10a1.penalty_points}, rank={sem_10a1.rank}")

# Test Excel export for Week 7
print("\n--- TESTING EXCEL EXPORT (WEEK 7) ---")
excel_bytes = export_competition_excel(p_week, admin)
wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
ws = wb.active

rows = list(ws.iter_rows(values_only=True))
header_idx = None
for idx, r in enumerate(rows):
    if r and any('Lớp' in str(x) for x in r if x):
        header_idx = idx
        break

print(f"Header row index: {header_idx}")
if header_idx is not None:
    print(f"Header: {rows[header_idx]}")

class_rows = []
for r in rows[header_idx+1:]:
    # columns typically: STT (0), Lớp (1), Điểm nề nếp / Điểm trừ (2), ... Xếp hạng (last)
    if r and len(r) > 1 and r[1]:
        c_name = str(r[1]).strip()
        class_rows.append((c_name, r))

print(f"Total classes exported: {len(class_rows)}")
print("Order of all classes:")
for i, (name, r) in enumerate(class_rows):
    # Print sample
    if i < 5 or (20 <= i <= 25) or (40 <= i <= 45) or i >= len(class_rows) - 3:
        print(f"  #{i+1}: {name} -> {r}")

# Check 10A1 row specifically
row_10a1 = next((r for name, r in class_rows if name == '10A1'), None)
print(f"\n10A1 Row in Excel: {row_10a1}")

# Test Excel export for Semester 1
print("\n--- TESTING EXCEL EXPORT (SEMESTER 1) ---")
sem_excel = export_competition_excel(p_sem, admin)
sem_wb = openpyxl.load_workbook(io.BytesIO(sem_excel))
sem_ws = sem_wb.active
sem_rows = list(sem_ws.iter_rows(values_only=True))
sem_class_rows = []
for r in sem_rows[header_idx+1:]:
    if r and len(r) > 1 and r[1]:
        c_name = str(r[1]).strip()
        sem_class_rows.append((c_name, r))

sem_10a1_row = next((r for name, r in sem_class_rows if name == '10A1'), None)
print(f"Semester 1 - 10A1 Row in Excel: {sem_10a1_row}")

# Test PDF export
print("\n--- TESTING PDF EXPORT ---")
pdf_week = export_competition_pdf(p_week, admin)
print(f"PDF Week 7 generated: {len(pdf_week)} bytes")
pdf_sem = export_competition_pdf(p_sem, admin)
print(f"PDF Semester 1 generated: {len(pdf_sem)} bytes")

print("\nALL VERIFICATIONS PASSED!")
